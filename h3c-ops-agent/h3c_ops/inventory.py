#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""设备清单读取：支持 .xlsx / .csv / .json，中英文列名均可。"""
from __future__ import annotations

import csv
import ipaddress
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple

HEADER_MAP = {
    # 管理地址
    "ip": "ip", "ip地址": "ip", "管理ip": "ip", "地址": "ip", "host": "ip", "hostname_ip": "ip",
    # 设备名
    "设备名称": "name", "名称": "name", "name": "name", "主机名": "name",
    "sysname": "name", "设备名": "name",
    # 凭据
    "用户名": "username", "username": "username", "user": "username",
    "账号": "username", "用户": "username",
    "密码": "password", "password": "password", "口令": "password",
    "端口": "port", "port": "port", "ssh端口": "port",
    "super密码": "secret", "secret": "secret", "enable密码": "secret", "提权密码": "secret",
    # 元信息
    "分组": "group", "group": "group", "设备分组": "group", "业务": "group",
    "模板": "template", "template": "template", "zabbix模板": "template",
    "团体名": "community", "community": "community", "snmp团体名": "community",
    "位置": "location", "location": "location",
    "备注": "remark", "remark": "remark", "说明": "remark",
    "启用": "enabled", "enabled": "enabled", "是否启用": "enabled",
}

TEMPLATE_HEADER = [
    "设备名称", "IP", "用户名", "密码", "端口",
    "super密码", "分组", "Zabbix模板", "团体名", "位置", "备注", "启用",
]

HOSTNAME_RE = re.compile(
    r"^(?=.{1,253}$)([A-Za-z0-9]([A-Za-z0-9-]{0,61}[A-Za-z0-9])?)"
    r"(\.[A-Za-z0-9]([A-Za-z0-9-]{0,61}[A-Za-z0-9])?)*$"
)


@dataclass
class Device:
    row: int
    ip: str
    username: str
    password: str = field(repr=False)
    port: int = 22
    name: str = ""
    secret: str = field(default="", repr=False)
    group: str = ""
    template: str = ""
    community: str = ""
    location: str = ""
    remark: str = ""
    enabled: bool = True
    credential_source: str = ""  # 凭据来自哪里：设备清单 / 配置默认 / 命令行

    @property
    def label(self) -> str:
        return self.name or self.ip

    @property
    def target(self) -> str:
        return f"{self.label}({self.ip}:{self.port})"


@dataclass
class Rejected:
    row: int
    ip: str
    reason: str


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value).strip()


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", "", _cell_to_str(value)).lower()


def valid_host(text: str) -> bool:
    try:
        ipaddress.ip_address(text)
        return True
    except ValueError:
        if re.fullmatch(r"[\d.]+", text):  # 形如 IPv4 但非法，如 300.1.1.1
            return False
        return bool(HOSTNAME_RE.match(text))


def _truthy(value: str) -> bool:
    return value.strip().lower() not in {"否", "no", "n", "0", "false", "off", "禁用", "停用"}


def _rows_from_excel(path: Path) -> List[Sequence[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return [row for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _rows_from_csv(path: Path) -> List[Sequence[Any]]:
    for encoding in ("utf-8-sig", "gbk"):
        try:
            with path.open("r", newline="", encoding=encoding) as f:
                return [tuple(row) for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"CSV 编码无法识别（试过 utf-8 / gbk）: {path}")


def _rows_from_json(path: Path) -> List[Sequence[Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        data = data.get("devices", [])
    if not isinstance(data, list) or not data:
        raise ValueError(f"JSON 清单应为非空数组: {path}")
    keys: List[str] = []
    for item in data:
        for key in item:
            if key not in keys:
                keys.append(key)
    return [keys] + [[item.get(k) for k in keys] for item in data]


def read_inventory(
    path: Path, default_user: str = "", default_password: str = "", default_secret: str = ""
) -> Tuple[List[Device], List[Rejected]]:
    """读取清单，返回 (有效设备, 被拒绝的行)。"""
    if not path.exists():
        raise FileNotFoundError(f"设备清单不存在: {path}")

    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows = _rows_from_excel(path)
    elif suffix in (".csv", ".txt"):
        rows = _rows_from_csv(path)
    elif suffix == ".json":
        rows = _rows_from_json(path)
    else:
        raise ValueError(f"不支持的清单格式 {suffix}，请用 .xlsx / .csv / .json")

    if not rows:
        raise ValueError("设备清单为空")

    header = rows[0]
    cols: Dict[str, int] = {}
    for index, raw in enumerate(header):
        key = HEADER_MAP.get(_norm_header(raw))
        if key and key not in cols:
            cols[key] = index

    if "ip" not in cols:
        raise ValueError(
            f"设备清单缺少 IP 列，当前表头: {[_cell_to_str(h) for h in header]}"
        )

    devices: List[Device] = []
    rejected: List[Rejected] = []
    seen: Dict[Tuple[str, int], int] = {}

    for line_no, row in enumerate(rows[1:], start=2):
        if not any(_cell_to_str(c) for c in row):
            continue

        def get(key: str) -> str:
            index = cols.get(key)
            if index is None or index >= len(row):
                return ""
            return _cell_to_str(row[index])

        ip = get("ip")
        port_text = get("port") or "22"
        row_user, row_password, row_secret = get("username"), get("password"), get("secret")
        device = Device(
            row=line_no,
            ip=ip,
            username=row_user or default_user,
            password=row_password or default_password,
            secret=row_secret or default_secret,
            name=get("name"),
            group=get("group"),
            template=get("template"),
            community=get("community"),
            location=get("location"),
            remark=get("remark"),
            enabled=_truthy(get("enabled") or "是"),
        )
        if row_user and row_password:
            device.credential_source = "设备清单"
        elif row_user or row_password:
            device.credential_source = "清单+配置默认"
        else:
            device.credential_source = "配置默认"

        problems: List[str] = []
        if not ip:
            problems.append("IP 为空")
        elif not valid_host(ip):
            problems.append(f"IP/域名格式不合法: {ip!r}")
        if not device.username:
            problems.append("用户名为空（也没有配置默认凭据）")
        if not device.password:
            problems.append("密码为空（也没有配置默认凭据）")
        try:
            device.port = int(float(port_text))
            if not 1 <= device.port <= 65535:
                raise ValueError
        except (TypeError, ValueError):
            problems.append(f"端口不合法: {port_text!r}")

        if not problems:
            key = (ip.lower(), device.port)
            if key in seen:
                problems.append(f"与第 {seen[key]} 行重复")
            else:
                seen[key] = line_no

        if problems:
            rejected.append(Rejected(line_no, ip or "-", "；".join(problems)))
        else:
            devices.append(device)

    return devices, rejected


def write_template(path: Path) -> Path:
    """生成设备清单模板（xlsx）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "devices"
    ws.append(TEMPLATE_HEADER)
    ws.append([
        "SW-Core-01", "10.12.20.1", "netops", "YourPassword", 22,
        "", "核心区", "", "OpsRead@2026", "DC-A-3F", "示例行，可直接删除", "是",
    ])
    ws.append([
        "SW-Acc-01", "10.12.21.11", "", "", 22,
        "", "接入区", "", "", "DC-A-2F", "留空则用 config.json 默认凭据", "是",
    ])

    head_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [18, 16, 12, 16, 8, 12, 12, 22, 18, 14, 28, 10]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    # 密码/端口列设为文本，避免 Excel 吃掉前导 0
    for col in (4, 5):
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "@"
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
