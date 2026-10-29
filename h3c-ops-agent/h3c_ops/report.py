#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""报告输出：控制台摘要 + Markdown + Excel。"""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from .inspect import LEVEL_CRITICAL, LEVEL_INFO, LEVEL_WARNING
from .models import DeviceResult

LEVEL_ORDER = {LEVEL_CRITICAL: 0, LEVEL_WARNING: 1, LEVEL_INFO: 2}


def summarize(results: List[DeviceResult]) -> Dict[str, Any]:
    levels = Counter()
    for result in results:
        levels[result.level] += 1

    issues = []
    for result in results:
        if result.inspection:
            for issue in result.inspection.issues:
                if issue.level in (LEVEL_CRITICAL, LEVEL_WARNING):
                    issues.append((issue.level, result.label, result.dev.ip, issue.item, issue.detail))
        if not result.reachable:
            issues.append((LEVEL_CRITICAL, result.label, result.dev.ip, "登录", result.error or "无法登录"))
    issues.sort(key=lambda item: (LEVEL_ORDER.get(item[0], 9), item[1]))

    def stage_stat(field_name: str) -> Counter:
        counter: Counter = Counter()
        for result in results:
            stage = getattr(result, field_name)
            if stage is not None:
                status = stage.status if hasattr(stage, "status") else "未知"
                counter[status] += 1
        return counter

    reachable = [r for r in results if r.reachable]
    scores = [r.score for r in reachable]
    return {
        "total": len(results),
        "reachable": len(reachable),
        "unreachable": len(results) - len(reachable),
        "levels": dict(levels),
        "avg_score": round(sum(scores) / len(scores), 1) if scores else 0,
        "issues": issues,
        "snmp": stage_stat("snmp"),
        "zabbix": stage_stat("zabbix"),
        "backup": stage_stat("backup"),
        "backup_changed": sum(1 for r in results if r.backup and r.backup.changed),
    }


def _counter_text(counter: Counter, order: List[str]) -> str:
    parts = [f"{key} {counter[key]}" for key in order if counter.get(key)]
    parts += [f"{key} {value}" for key, value in counter.items() if key not in order]
    return " / ".join(parts) if parts else "-"


def console_summary(results: List[DeviceResult], meta: Dict[str, Any]) -> str:
    stats = summarize(results)
    lines = []
    lines.append("")
    lines.append("=" * 78)
    lines.append(f"华三交换机批处理完成  {meta.get('finished_at', '')}   模式: {meta.get('mode', '-')}")
    lines.append("=" * 78)
    lines.append(
        f"设备 {stats['total']} 台 | 可达 {stats['reachable']} | 不可达 {stats['unreachable']} | "
        f"平均健康分 {stats['avg_score']}"
    )
    level_text = " ".join(
        f"{name} {stats['levels'].get(name, 0)}" for name in ("正常", "关注", "警告", "严重", "不可达")
        if stats["levels"].get(name)
    )
    if level_text:
        lines.append(f"健康分布: {level_text}")
    if stats["snmp"]:
        lines.append("团体名: " + _counter_text(stats["snmp"], ["成功", "无变化", "预览", "失败"]))
    if stats["zabbix"]:
        lines.append("Zabbix: " + _counter_text(stats["zabbix"], ["成功", "无变化", "预览", "冲突", "失败"]))
    if stats["backup"]:
        lines.append("备份: " + _counter_text(stats["backup"], ["成功", "有变更", "失败"])
                     + f"（配置有变更 {stats['backup_changed']} 台）")

    if stats["issues"]:
        lines.append("")
        lines.append("需要处理：")
        for index, (level, label, ip, item, detail) in enumerate(stats["issues"][:15], start=1):
            lines.append(f"  {index:>2}. [{level}] {label}({ip}) {item}: {detail}")
        if len(stats["issues"]) > 15:
            lines.append(f"  ... 其余 {len(stats['issues']) - 15} 条见报告")
    else:
        lines.append("")
        lines.append("未发现需要处理的告警项。")
    lines.append("")
    return "\n".join(lines)


def write_markdown(results: List[DeviceResult], path: Path, meta: Dict[str, Any]) -> Path:
    stats = summarize(results)
    lines: List[str] = []
    lines.append("# 华三交换机批量巡检 / 接入报告")
    lines.append("")
    lines.append(f"- 生成时间：{meta.get('finished_at', '')}")
    lines.append(f"- 运行模式：{meta.get('mode', '-')}")
    lines.append(f"- 执行阶段：{' → '.join(meta.get('stages', [])) or '-'}")
    lines.append(f"- 设备总数：{stats['total']}（可达 {stats['reachable']}，不可达 {stats['unreachable']}）")
    lines.append(f"- 平均健康分：{stats['avg_score'] if stats['reachable'] else '-'}")
    lines.append("")

    lines.append("## 总体结论")
    lines.append("")
    if stats["issues"]:
        for level, label, ip, item, detail in stats["issues"][:10]:
            lines.append(f"- **[{level}]** {label}（{ip}）{item}：{detail}")
        if len(stats["issues"]) > 10:
            lines.append(f"- 其余 {len(stats['issues']) - 10} 条见下方「需要处理的问题」")
    else:
        lines.append("- 所有设备指标正常，未发现需要处理的告警项。")
    if stats["snmp"]:
        lines.append(f"- 团体名：{_counter_text(stats['snmp'], ['成功', '无变化', '预览', '失败'])}")
    if stats["zabbix"]:
        lines.append(f"- Zabbix 接入：{_counter_text(stats['zabbix'], ['成功', '无变化', '预览', '冲突', '失败'])}")
    if stats["backup"]:
        lines.append(f"- 配置备份：{_counter_text(stats['backup'], ['成功', '有变更', '失败'])}，其中配置发生变更 {stats['backup_changed']} 台")
    lines.append("")

    if stats["issues"]:
        lines.append("## 需要处理的问题")
        lines.append("")
        lines.append("| 等级 | 设备 | IP | 项目 | 说明 |")
        lines.append("|------|------|----|------|------|")
        for level, label, ip, item, detail in stats["issues"]:
            lines.append(f"| {level} | {label} | {ip} | {item} | {detail} |")
        lines.append("")

    lines.append("## 逐台明细")
    lines.append("")
    lines.append("| 设备 | 管理IP | 型号 | 运行时长 | CPU | 内存 | 温度 | 端口up/总 | 健康分 | 等级 |")
    lines.append("|------|--------|------|----------|-----|------|------|-----------|--------|------|")
    for result in results:
        insp = result.inspection
        if not insp:
            lines.append(
                f"| {result.label} | {result.dev.ip} | - | - | - | - | - | - | 0 | {result.level} |"
            )
            continue
        uptime = f"{insp.uptime_days:.0f} 天" if insp.uptime_days is not None else "-"
        cpu = f"{insp.cpu}%" if insp.cpu is not None else "-"
        mem = f"{insp.mem_percent:.0f}%" if insp.mem_percent is not None else "-"
        temp = f"{insp.max_temp:.0f}℃" if insp.max_temp is not None else "-"
        lines.append(
            f"| {insp.hostname or result.dev.label} | {result.dev.ip} | {insp.model or '-'} | {uptime} | "
            f"{cpu} | {mem} | {temp} | {insp.if_up}/{insp.if_total} | {insp.score} | {insp.level} |"
        )
    lines.append("")

    lines.append("## 执行结果")
    lines.append("")
    lines.append("| 设备 | IP | 团体名 | Zabbix | 配置备份 | 备份文件 |")
    lines.append("|------|----|--------|--------|----------|----------|")
    for result in results:
        snmp = result.snmp.status if result.snmp else "-"
        zabbix = result.zabbix.status if result.zabbix else "-"
        backup = result.backup.status if result.backup else "-"
        backup_path = ""
        if result.backup and result.backup.path:
            backup_path = Path(result.backup.path).name
        lines.append(
            f"| {result.label} | {result.dev.ip} | {snmp} | {zabbix} | {backup} | {backup_path or '-'} |"
        )
    lines.append("")

    warnings = []
    for result in results:
        warnings.extend(f"{result.label}: {w}" for w in result.warnings)
        for stage in (result.snmp, result.zabbix):
            if stage and stage.warnings:
                warnings.extend(f"{result.label}: {w}" for w in stage.warnings)
    if warnings:
        lines.append("## 需要注意")
        lines.append("")
        for item in warnings:
            lines.append(f"- {item}")
        lines.append("")

    lines.append("## 建议动作")
    lines.append("")
    suggestions = []
    if stats["unreachable"]:
        suggestions.append("不可达设备先确认管理网路由、SSH 是否开启、账号是否被锁。")
    if stats["levels"].get("严重"):
        suggestions.append("健康度为「严重」的设备优先处理（风扇/电源/温度/CPU 类问题影响面最大）。")
    if stats["backup_changed"]:
        suggestions.append("配置有变更的设备，建议核对变更单（.diff 文件在同目录）。")
    suggestions.append("Zabbix 侧建议观察 24 小时，确认 SNMP 数据正常采集后再开启触发器告警。")
    for item in suggestions:
        lines.append(f"- {item}")
    lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_xlsx(results: List[DeviceResult], path: Path, meta: Dict[str, Any]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    stats = summarize(results)
    wb = Workbook()

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    level_fill = {
        "正常": PatternFill("solid", fgColor="E2EFDA"),
        "关注": PatternFill("solid", fgColor="FFF2CC"),
        "警告": PatternFill("solid", fgColor="FCE4D6"),
        "严重": PatternFill("solid", fgColor="F8CBAD"),
        "不可达": PatternFill("solid", fgColor="D9D9D9"),
    }

    def style_sheet(ws, widths: List[int], freeze: str = "A2") -> None:
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 26

    # ---- 巡检汇总 ----
    ws = wb.active
    ws.title = "巡检汇总"
    ws.append([
        "序号", "设备名", "管理IP", "型号", "Comware", "运行时长(天)", "CPU%", "内存%",
        "最高温(℃)", "风扇", "电源", "端口up/总", "健康分", "等级", "SNMP", "Zabbix", "备份",
        "登录账号", "主要问题",
    ])
    for index, result in enumerate(results, start=1):
        insp = result.inspection
        if not insp:
            ws.append([
                index, result.label, result.dev.ip, "-", "-", None, None, None, None, "-", "-",
                "-", 0, result.level, result.snmp.status if result.snmp else "-",
                result.zabbix.status if result.zabbix else "-",
                result.backup.status if result.backup else "-",
                result.login_account or "-",
                result.error or "无法登录",
            ])
            continue
        ws.append([
            index,
            insp.hostname or result.dev.label,
            result.dev.ip,
            insp.model,
            f"{insp.comware} {insp.release}".strip() if insp.comware else "-",
            round(insp.uptime_days, 1) if insp.uptime_days is not None else None,
            insp.cpu,
            round(insp.mem_percent, 1) if insp.mem_percent is not None else None,
            round(insp.max_temp, 1) if insp.max_temp is not None else None,
            "异常" if insp.fan_bad else (f"{insp.fan_total} OK" if insp.fan_total else "-"),
            "异常" if insp.power_bad else (f"{insp.power_total} OK" if insp.power_total else "-"),
            f"{insp.if_up}/{insp.if_total}",
            insp.score,
            insp.level,
            result.snmp.status if result.snmp else "-",
            result.zabbix.status if result.zabbix else "-",
            result.backup.status if result.backup else "-",
            result.login_account or "-",
            "；".join(result.problems()[:3]),
        ])
    style_sheet(ws, [5, 18, 15, 20, 20, 12, 8, 8, 10, 10, 10, 11, 8, 8, 10, 12, 10, 14, 45])
    for row in ws.iter_rows(min_row=2, min_col=14, max_col=14):
        for cell in row:
            fill = level_fill.get(str(cell.value))
            if fill:
                cell.fill = fill

    # ---- 问题明细 ----
    ws = wb.create_sheet("问题明细")
    ws.append(["设备", "管理IP", "等级", "项目", "说明"])
    for result in results:
        if not result.reachable:
            ws.append([result.label, result.dev.ip, LEVEL_CRITICAL, "登录", result.error or "无法登录"])
            continue
        if not result.inspection:
            continue
        for issue in sorted(result.inspection.issues, key=lambda i: LEVEL_ORDER.get(i.level, 9)):
            ws.append([result.label, result.dev.ip, issue.level, issue.item, issue.detail])
    style_sheet(ws, [18, 15, 8, 12, 70])

    # ---- 执行结果 ----
    ws = wb.create_sheet("执行结果")
    ws.append(["设备", "管理IP", "阶段", "结果", "说明"])
    for result in results:
        for stage in result.all_stages():
            ws.append([result.label, result.dev.ip, stage.name, stage.status, stage.detail])
    style_sheet(ws, [18, 15, 12, 10, 70])

    # ---- 运行信息 ----
    ws = wb.create_sheet("运行信息")
    ws.append(["项目", "内容"])
    rows = [
        ("生成时间", meta.get("finished_at", "")),
        ("开始时间", meta.get("started_at", "")),
        ("运行模式", meta.get("mode", "")),
        ("执行阶段", " → ".join(meta.get("stages", []))),
        ("设备总数", stats["total"]),
        ("可达 / 不可达", f"{stats['reachable']} / {stats['unreachable']}"),
        ("平均健康分", stats["avg_score"]),
        ("告警数量", len(stats["issues"])),
        ("配置文件", meta.get("config_file", "")),
        ("设备清单", meta.get("inventory_file", "")),
        ("Zabbix 地址", meta.get("zabbix_url", "-")),
    ]
    for key, value in rows:
        ws.append([key, value])
    style_sheet(ws, [18, 80], freeze="A2")
    for cell in ws["A"]:
        cell.font = Font(bold=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
