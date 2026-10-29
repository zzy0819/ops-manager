#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""报告输出单测：控制台摘要 / Markdown / Excel 都能正常生成。"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from h3c_ops.backup import BackupResult
from h3c_ops.inspect import Inspection, Issue
from h3c_ops.inventory import Device
from h3c_ops.models import DeviceResult, Stage
from h3c_ops.report import console_summary, summarize, write_markdown, write_xlsx


def build_results():
    healthy = DeviceResult(
        dev=Device(row=2, ip="10.12.20.1", username="u", password="p", name="SW-Core-01"),
        reachable=True,
        inspection=Inspection(hostname="SW-Core-01", model="S5560X-30C-EI", cpu=12,
                              mem_percent=45.0, max_temp=40, if_total=48, if_up=48,
                              score=100, level="正常",
                              snmp={"has_snmp": True, "v2c": True}),
        snmp=Stage("团体名", "成功", "团体名已生效"),
        zabbix=Stage("Zabbix接入", "成功", "新建主机"),
        backup=BackupResult("有变更", "配置有变更：+2 / -1 行", "/tmp/x/20260101.cfg", True, 2, 1),
    )
    broken = DeviceResult(
        dev=Device(row=3, ip="10.12.20.2", username="u", password="p", name="SW-Acc-02"),
        reachable=True,
        inspection=Inspection(hostname="SW-Acc-02", model="S5130S-28P-EI", cpu=91,
                              mem_percent=90.0, max_temp=72, fan_bad=["Fan2(Fault)"],
                              if_total=28, if_up=20, down_ports=["GE1/0/5"],
                              score=42, level="严重",
                              issues=[Issue("严重", "CPU", "91% 超过阈值 80%")]),
        snmp=Stage("团体名", "失败", "命令执行失败"),
        zabbix=Stage("Zabbix接入", "失败", "无法连接 Zabbix"),
        backup=BackupResult("失败", "抓取配置失败"),
    )
    unreachable = DeviceResult(
        dev=Device(row=4, ip="10.12.20.3", username="u", password="p", name="SW-Acc-03"),
        reachable=False,
        error="连接超时（15s）",
        snmp=Stage("团体名", "失败", "设备不可达，已跳过"),
        backup=BackupResult("失败", "设备不可达，已跳过"),
    )
    return [healthy, broken, unreachable]


class ReportTest(unittest.TestCase):
    def setUp(self):
        self.results = build_results()
        self.meta = {
            "started_at": "2026-10-29 22:00:00",
            "finished_at": "2026-10-29 22:05:00",
            "mode": "实际执行",
            "stages": ["巡检", "团体名", "Zabbix接入", "配置备份"],
            "config_file": "config.json",
            "inventory_file": "devices.xlsx",
            "zabbix_url": "http://10.12.10.50",
        }
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_summarize(self):
        stats = summarize(self.results)
        self.assertEqual(stats["total"], 3)
        self.assertEqual(stats["reachable"], 2)
        self.assertEqual(stats["unreachable"], 1)
        self.assertEqual(stats["levels"]["正常"], 1)
        self.assertEqual(stats["levels"]["不可达"], 1)
        self.assertEqual(stats["backup_changed"], 1)
        self.assertTrue(any("登录" == item[3] for item in stats["issues"]))

    def test_console_summary_mentions_key_numbers(self):
        text = console_summary(self.results, self.meta)
        self.assertIn("设备 3 台", text)
        self.assertIn("不可达 1", text)
        self.assertIn("需要处理", text)

    def test_markdown(self):
        path = write_markdown(self.results, self.root / "summary.md", self.meta)
        content = path.read_text(encoding="utf-8")
        self.assertIn("# 华三交换机批量巡检 / 接入报告", content)
        self.assertIn("## 总体结论", content)
        self.assertIn("## 需要处理的问题", content)
        self.assertIn("SW-Core-01", content)
        self.assertIn("不可达", content)

    def test_xlsx(self):
        path = write_xlsx(self.results, self.root / "report.xlsx", self.meta)
        self.assertTrue(path.exists())
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["巡检汇总", "问题明细", "执行结果", "运行信息"])
        ws = wb["巡检汇总"]
        self.assertEqual(ws.max_row, 4)  # 表头 + 3 台
        self.assertEqual(ws.cell(row=2, column=2).value, "SW-Core-01")
        self.assertEqual(ws.cell(row=3, column=14).value, "严重")
        self.assertIn("运行信息", wb.sheetnames)
        wb.close()

    def test_xlsx_handles_missing_inspection(self):
        path = write_xlsx([self.results[2]], self.root / "one.xlsx", self.meta)
        wb = load_workbook(path)
        self.assertEqual(wb["巡检汇总"].cell(row=2, column=2).value, "SW-Acc-03")
        wb.close()


if __name__ == "__main__":
    unittest.main()
