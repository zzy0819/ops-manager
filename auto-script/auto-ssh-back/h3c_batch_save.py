#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
批量 SSH 登录华三(H3C/Comware)交换机并保存配置（增强健壮性版）

依赖:
    pip install netmiko openpyxl

Excel 格式（第一行为表头，中英文列名均可）:
    IP | 用户名 | 密码 | 端口(可选,默认22) | 设备名称(可选) | super密码(可选)
    * 密码如果是纯数字，请把单元格设置为「文本」格式，否则 Excel 会吃掉前导 0。

用法:
    python h3c_batch_save.py --template                  # 生成模板
    python h3c_batch_save.py -f devices.xlsx --check-only  # 只校验端口+账号密码，不保存
    python h3c_batch_save.py -f devices.xlsx             # 校验通过后保存配置
    python h3c_batch_save.py -f devices.xlsx --backup    # 先把配置备份到 ./backup/*.cfg，再保存
    python h3c_batch_save.py -f devices.xlsx -w 10 -t 15 -r 2

处理流程（每台设备）:
    1. 参数校验   : IP/域名、端口范围、用户名密码非空、重复项
    2. 端口检测   : TCP 连通性 + 读取 SSH banner，确认确实是 SSH 服务
    3. 登录认证   : 认证失败【不重试】，避免触发设备账号锁定；网络类错误才重试
    4. 登录后校验 : 识别提示符、退出系统视图、检查是否 Comware 设备
    5. 保存配置   : save force，并根据回显判断成功/失败/无法确认
退出码: 0=全部成功, 1=存在失败/未确认, 130=用户中断
"""
from __future__ import annotations

import argparse
import csv
import ipaddress
import logging
import re
import socket
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple

try:
    import paramiko
    from netmiko import ConnectHandler
    from netmiko.exceptions import (
        NetmikoAuthenticationException,
        NetmikoTimeoutException,
    )
    try:  # netmiko >= 4.x
        from netmiko.exceptions import ReadTimeout
    except ImportError:  # 老版本没有该异常
        class ReadTimeout(Exception):
            pass
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError as e:
    sys.exit(f"缺少依赖: {e}\n请先执行: pip install netmiko openpyxl")

log = logging.getLogger("h3c_save")

# ---------------------------------------------------------------- 常量 ----
HEADER_MAP = {
    "ip": "ip", "ip地址": "ip", "host": "ip", "管理ip": "ip", "地址": "ip",
    "用户名": "username", "username": "username", "user": "username",
    "账号": "username", "用户": "username",
    "密码": "password", "password": "password",
    "端口": "port", "port": "port",
    "设备名称": "name", "名称": "name", "name": "name", "hostname": "name",
    "super密码": "secret", "super": "secret", "secret": "secret", "enable密码": "secret",
}

HOSTNAME_RE = re.compile(
    r"^(?=.{1,253}$)([A-Za-z0-9]([A-Za-z0-9-]{0,61}[A-Za-z0-9])?)(\.[A-Za-z0-9]([A-Za-z0-9-]{0,61}[A-Za-z0-9])?)*$"
)
# Comware 提示符：用户视图 <SW1>，系统视图 [SW1] / [SW1-GigabitEthernet1/0/1]
PROMPT_RE = re.compile(r"^[<\[](?P<name>[^<>\[\]]+)[>\]]$")

SAVE_FAIL_KEYWORDS = (
    "error", "invalid", "unrecognized", "failed", "permission denied",
    "incomplete command", "too many parameters", "wrong parameter",
    "not enough", "no space", "not permitted", "authorization failed",
)
SAVE_OK_KEYWORDS = ("successfully", "saved the current configuration", "save the configuration successfully")

NEGOTIATION_MARKERS = (
    "no acceptable", "incompatible ssh peer", "no matching",
    "unable to negotiate", "no common",
)

_legacy_lock = threading.Lock()
_legacy_enabled = False


def is_negotiation_error(e: Exception) -> bool:
    """是否为 SSH 算法协商失败（kex/cipher/mac/hostkey 与设备无交集）。"""
    s = str(e).lower()
    return any(k in s for k in NEGOTIATION_MARKERS)


def _paramiko_major() -> int:
    try:
        return int(str(paramiko.__version__).split(".")[0])
    except Exception:  # noqa: BLE001
        return 0


def enable_legacy_algorithms() -> bool:
    """
    给 paramiko 追加老设备常用的旧算法（追加在列表末尾，新设备仍优先协商强算法）。
    包括 diffie-hellman-group1/14-sha1、group-exchange-sha1、CBC 类加密、ssh-rsa/ssh-dss。
    返回 True 表示本次是首次启用。全局生效，线程安全，可重复调用。
    """
    global _legacy_enabled
    with _legacy_lock:
        if _legacy_enabled:
            return False
        T = paramiko.Transport
        extras = {
            "_preferred_kex": ("_kex_info", (
                "diffie-hellman-group14-sha1",
                "diffie-hellman-group-exchange-sha1",
                "diffie-hellman-group1-sha1",
            )),
            "_preferred_ciphers": ("_cipher_info", (
                "aes128-cbc", "aes192-cbc", "aes256-cbc", "3des-cbc",
            )),
            "_preferred_macs": ("_mac_info", ("hmac-sha1", "hmac-md5")),
            "_preferred_keys": ("_key_info", ("ssh-rsa", "ssh-dss")),
        }
        for pref_attr, (info_attr, wanted) in extras.items():
            cur = tuple(getattr(T, pref_attr, ()))
            info = getattr(T, info_attr, {})
            add = tuple(x for x in wanted if x in info and x not in cur)
            if add:
                setattr(T, pref_attr, cur + add)
        _legacy_enabled = True
        return True


STATUS_OK, STATUS_CHECK_OK = "成功", "校验通过"
STATUS_FAIL, STATUS_SKIP, STATUS_UNSURE = "失败", "跳过", "未确认"


# ---------------------------------------------------------------- 数据 ----
@dataclass
class Device:
    row: int
    ip: str
    username: str
    password: str = field(repr=False)  # repr 中隐藏密码，避免日志泄露
    port: int = 22
    name: str = ""
    secret: str = field(default="", repr=False)


@dataclass
class Result:
    dev: Device
    status: str
    stage: str = ""
    detail: str = ""
    prompt: str = ""
    elapsed: float = 0.0


# ---------------------------------------------------------- Excel 读取 ----
def make_template(path: str = "devices_template.xlsx") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "devices"
    ws.append(["IP", "用户名", "密码", "端口", "设备名称", "super密码"])
    ws.append(["192.168.1.1", "admin", "your_password", 22, "SW-Core-01", ""])
    wb.save(path)
    print(f"模板已生成: {path}")


def _cell_to_str(v) -> str:
    """把单元格值安全转成字符串（处理数字 22.0、前后空格等）。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _norm_header(h) -> str:
    return re.sub(r"\s+", "", _cell_to_str(h)).lower()


def read_devices(path: str) -> Tuple[List[Device], List[Result]]:
    """读取 Excel，返回 (有效设备, 因参数问题被直接判定的结果)。"""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise SystemExit(f"文件不存在: {path}")
    except (InvalidFileException, KeyError, OSError) as e:
        raise SystemExit(f"无法打开 Excel（仅支持 .xlsx/.xlsm，.xls 请先另存为 .xlsx）: {e}")

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise SystemExit("Excel 为空")

    cols = {}
    for idx, h in enumerate(header):
        key = HEADER_MAP.get(_norm_header(h))
        if key and key not in cols:
            cols[key] = idx
    missing = [k for k in ("ip", "username", "password") if k not in cols]
    if missing:
        raise SystemExit(f"Excel 缺少必需列: {missing}，当前表头: {[_cell_to_str(h) for h in header]}")

    def get(row, key) -> str:
        i = cols.get(key)
        return _cell_to_str(row[i]) if i is not None and i < len(row) else ""

    devices: List[Device] = []
    rejected: List[Result] = []
    seen = {}

    for n, row in enumerate(rows, start=2):
        if not any(_cell_to_str(c) for c in row):
            continue  # 完全空行
        ip = get(row, "ip")
        port_s = get(row, "port") or "22"
        dev = Device(
            row=n, ip=ip, username=get(row, "username"), password=get(row, "password"),
            name=get(row, "name") or ip, secret=get(row, "secret"),
        )

        errors = []
        if not ip:
            errors.append("IP 为空")
        elif not _valid_host(ip):
            errors.append(f"IP/域名格式不合法: {ip!r}")
        if not dev.username:
            errors.append("用户名为空")
        if not dev.password:
            errors.append("密码为空")
        try:
            dev.port = int(port_s)
            if not 1 <= dev.port <= 65535:
                raise ValueError
        except ValueError:
            errors.append(f"端口不合法: {port_s!r}")

        if not errors:
            key = (ip.lower(), dev.port)
            if key in seen:
                errors.append(f"与第 {seen[key]} 行重复")
            else:
                seen[key] = n

        if errors:
            rejected.append(Result(dev, STATUS_SKIP, "参数校验", "；".join(errors)))
        else:
            devices.append(dev)

    wb.close()
    return devices, rejected


def _valid_host(s: str) -> bool:
    try:
        ipaddress.ip_address(s)
        return True
    except ValueError:
        if re.fullmatch(r"[\d.]+", s):  # 形如 IPv4 但不合法，如 300.1.1.1 / 10.0.0
            return False
        return bool(HOSTNAME_RE.match(s))


# ---------------------------------------------------------- 端口检测 ----
def check_ssh_port(host: str, port: int, timeout: float, retries: int) -> Tuple[bool, str]:
    """检测 TCP 端口是否开放，并读取 banner 确认为 SSH 服务。"""
    last_err = ""
    for attempt in range(1, retries + 2):
        try:
            with socket.create_connection((host, port), timeout=timeout) as s:
                s.settimeout(min(timeout, 5))
                try:
                    banner = s.recv(255)
                except socket.timeout:
                    return True, "端口开放（未读到 banner，继续尝试）"
                if banner and not banner.startswith(b"SSH-"):
                    return False, f"端口开放但不是 SSH 服务（banner: {banner[:30]!r}）"
                if not banner:
                    return False, "端口开放但对端立即关闭连接"
                return True, banner.decode(errors="ignore").strip()
        except ConnectionRefusedError:
            return False, f"端口 {port} 未开放/连接被拒绝（SSH 可能未启用）"  # 不重试
        except socket.gaierror as e:
            return False, f"域名解析失败: {e}"  # 不重试
        except (socket.timeout, TimeoutError):
            last_err = f"TCP 连接超时（{timeout}s）"
        except OSError as e:
            last_err = f"网络错误: {e}"
        if attempt <= retries:
            time.sleep(min(2 * attempt, 5))
    return False, f"{last_err}，已重试 {retries} 次，请检查路由/ACL/防火墙"


# ---------------------------------------------------------- 登录与保存 ----
def _connect(dev: Device, timeout: int, retries: int):
    """登录设备。认证失败立即抛出（不重试，防止账号被锁），网络类错误按次数重试。"""
    params = dict(
        device_type="hp_comware", host=dev.ip, port=dev.port,
        username=dev.username, password=dev.password,
        conn_timeout=timeout, banner_timeout=timeout, auth_timeout=timeout,
        fast_cli=False,
        # 不使用公钥/agent，只做账号密码认证，避免因本机密钥过多触发服务端限制
        use_keys=False, allow_agent=False,
    )
    attempt = 0
    tried_legacy = False
    while True:
        attempt += 1
        try:
            return ConnectHandler(**params)
        except NetmikoAuthenticationException:
            raise
        except Exception as e:  # noqa: BLE001
            log.debug("%s 第 %d 次登录失败: %s", dev.ip, attempt, e)
            if is_negotiation_error(e):
                # 算法协商失败：重试没有意义，只尝试一次启用旧算法后重连
                if tried_legacy:
                    raise
                tried_legacy = True
                if enable_legacy_algorithms():
                    log.warning("检测到设备仅支持旧版 SSH 算法（%s），已启用兼容旧算法后重试。"
                                "建议在设备上启用更安全的算法。", dev.ip)
                attempt -= 1  # 不占用网络重试次数
                continue
            if attempt > retries:
                raise
            time.sleep(min(2 * attempt, 5))


def _output_has(text: str, words) -> bool:
    low = text.lower()
    return any(w in low for w in words)


def process_device(dev: Device, timeout: int, retries: int, check_only: bool,
                   backup_dir: Optional[str] = None) -> Result:
    t0 = time.time()

    def done(status, stage="", detail="", prompt="") -> Result:
        return Result(dev, status, stage, detail, prompt, round(time.time() - t0, 1))

    # 1) 端口检测
    ok, msg = check_ssh_port(dev.ip, dev.port, timeout, retries)
    if not ok:
        return done(STATUS_FAIL, "端口检测", msg)

    # 2) 登录认证
    try:
        conn = _connect(dev, timeout, retries)
    except NetmikoAuthenticationException:
        return done(STATUS_FAIL, "账号认证", "用户名或密码错误（已停止重试以避免账号锁定），或该账号无 SSH 登录权限")
    except Exception as e:  # noqa: BLE001
        # 注意：netmiko 会把 paramiko 的 SSHException 包装成 NetmikoTimeoutException，
        # 所以必须先按错误内容判断，不能只看异常类型
        if is_negotiation_error(e):
            if _paramiko_major() >= 5:
                return done(
                    STATUS_FAIL, "SSH 算法协商",
                    f"当前 paramiko {paramiko.__version__} 已移除 SHA-1 密钥交换和 ssh-rsa(SHA-1)，"
                    "无法连接只支持这些算法的老设备（如 Comware 5）。"
                    '请降级后重试: pip install "paramiko<5"。'
                    f"原始错误: {_short(e)}",
                )
            return done(
                STATUS_FAIL, "SSH 算法协商",
                "本机与设备没有共同的 SSH 算法（已尝试启用旧算法仍失败）。"
                "请用 `ssh -vv 用户@IP` 查看设备支持的算法，或在设备上执行 "
                "`ssh server key-exchange dh-group14-sha1`（命令因型号而异）。"
                f"原始错误: {_short(e)}",
            )
        if isinstance(e, NetmikoTimeoutException):
            return done(STATUS_FAIL, "SSH 登录", f"SSH 会话超时或被对端中断: {_short(e)}")
        return done(STATUS_FAIL, "SSH 登录", f"{type(e).__name__}: {_short(e)}")

    try:
        # 3) 登录后校验
        prompt = conn.find_prompt().strip()
        m = PROMPT_RE.match(prompt)
        if not m:
            return done(STATUS_FAIL, "登录后校验", f"无法识别设备提示符 {prompt!r}，可能不是 H3C 设备", prompt)

        if prompt.startswith("["):  # 位于系统视图，先退回用户视图
            conn.send_command_timing("return", read_timeout=10)
            prompt = conn.find_prompt().strip()

        note = ""
        try:
            ver = conn.send_command_timing("display version", read_timeout=30)
            if "comware" not in ver.lower() and "h3c" not in ver.lower():
                note = "；警告: display version 回显中未发现 Comware/H3C 字样"
        except ReadTimeout:
            note = "；警告: display version 超时"

        # 可选：把当前配置备份到本地（只读操作；失败则不继续保存，便于发现问题）
        if backup_dir:
            try:
                bpath, bwarn = backup_config(conn, dev, backup_dir)
                note += f"；已备份到 {bpath}{bwarn}"
            except Exception as e:  # noqa: BLE001
                return done(STATUS_FAIL, "备份配置", f"{_short(e)}（为安全起见未执行保存）", prompt)

        if check_only:
            return done(STATUS_CHECK_OK, "", f"端口开放、认证通过、提示符 {prompt}{note}", prompt)

        # 提权（可选）
        if dev.secret:
            out = conn.send_command_timing("super", read_timeout=15)
            if "assword" in out:
                out = conn.send_command_timing(dev.secret, read_timeout=15)
                if _output_has(out, ("incorrect", "invalid", "fail", "denied")):
                    return done(STATUS_FAIL, "提权", "super 密码错误", prompt)

        # 4) 保存配置
        out = conn.send_command_timing("save force", read_timeout=120, strip_prompt=False)
        # 兼容老版本：可能出现 [Y/N] 或文件名确认
        for _ in range(3):
            low = out.lower()
            if "[y/n]" in low or "(y/n)" in low:
                out += conn.send_command_timing("Y", read_timeout=120)
            elif "(to leave the existing filename" in low or "please input the file name" in low:
                out += conn.send_command_timing("\n", read_timeout=120)
            else:
                break

        brief = " ".join(out.split())[-200:]
        if _output_has(out, SAVE_FAIL_KEYWORDS):
            return done(STATUS_FAIL, "保存配置", f"设备返回错误: {brief}", prompt)
        if _output_has(out, SAVE_OK_KEYWORDS):
            return done(STATUS_OK, "", f"配置保存成功{note}", prompt)
        return done(STATUS_UNSURE, "保存配置", f"未在回显中确认保存结果，请人工核实: {brief}", prompt)

    except ReadTimeout:
        return done(STATUS_UNSURE, "保存配置", "等待设备回显超时（可能仍在保存，请人工核实）")
    except Exception as e:  # noqa: BLE001
        return done(STATUS_FAIL, "执行命令", f"{type(e).__name__}: {_short(e)}")
    finally:
        try:
            conn.disconnect()
        except Exception:  # noqa: BLE001
            pass


def _safe_filename(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_") or "device"


def backup_config(conn, dev: Device, backup_dir: str) -> Tuple[str, str]:
    """把设备当前运行配置读到本地 .cfg 文件。返回 (文件路径, 警告信息)；失败抛 RuntimeError。"""
    out = conn.send_command_timing("display current-configuration", read_timeout=180)
    text = re.sub(r"\x1b\[[0-9;]*[A-Za-z]", "", out)           # 去掉终端控制字符
    text = re.sub(r"[ \t]*-+\s*More\s*-+[ \t]*", "", text)       # 去掉分页残留
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(text) < 20 or _output_has(text[:300], ("unrecognized", "error", "permission denied",
                                                  "authorization failed", "incomplete command")):
        raise RuntimeError(f"读取配置失败: {' '.join(text.split())[:150]}")
    warn = ""
    if text.splitlines()[-1].strip().lower() != "return":      # Comware 配置以 return 结尾
        warn = "；警告: 配置末尾不是 return，文件可能不完整"
    folder = Path(backup_dir)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{_safe_filename(dev.name)}_{dev.ip}_{datetime.now():%Y%m%d_%H%M%S}.cfg"
    path.write_text(text + "\n", encoding="utf-8")
    return str(path), warn


def _short(e: Exception, n: int = 200) -> str:
    return " ".join(str(e).split())[:n]


# ---------------------------------------------------------------- 报告 ----
REPORT_HEADER = ["行号", "设备名称", "IP", "端口", "设备提示符", "结果", "失败阶段", "详情", "耗时(s)"]


def write_report(results: List[Result], stamp: str) -> str:
    rows = [
        [r.dev.row, r.dev.name, r.dev.ip, r.dev.port, r.prompt,
         r.status, r.stage, r.detail, r.elapsed]
        for r in sorted(results, key=lambda x: x.dev.row)
    ]
    path = f"save_config_result_{stamp}.xlsx"
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "result"
        ws.append(REPORT_HEADER)
        for r in rows:
            ws.append(r)
        wb.save(path)
    except Exception as e:  # noqa: BLE001  Excel 写失败（被占用/无权限）则降级为 CSV
        path = f"save_config_result_{stamp}.csv"
        log.warning("写 Excel 报告失败(%s)，改写 CSV: %s", e, path)
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(REPORT_HEADER)
            w.writerows(rows)
    return path


# ----------------------------------------------------------------- 主流程 ----
def main() -> int:
    ap = argparse.ArgumentParser(description="批量保存华三交换机配置")
    ap.add_argument("-f", "--file", help="设备清单 Excel 文件")
    ap.add_argument("-w", "--workers", type=int, default=5, help="并发数(1-50)，默认5")
    ap.add_argument("-t", "--timeout", type=int, default=20, help="连接/认证超时秒数，默认20")
    ap.add_argument("-r", "--retries", type=int, default=2, help="网络类错误重试次数，默认2（认证失败不重试）")
    ap.add_argument("--check-only", action="store_true", help="只校验端口与账号密码，不保存配置")
    ap.add_argument("--backup", nargs="?", const="backup", metavar="DIR",
                    help="把每台设备当前配置备份为本地 .cfg 文件，目录默认 ./backup（先备份再保存）")
    ap.add_argument("--legacy-ssh", action="store_true",
                    help="启动时直接启用旧版 SSH 算法（默认遇到协商失败时才自动启用）")
    ap.add_argument("--template", action="store_true", help="生成 Excel 模板并退出")
    ap.add_argument("-v", "--verbose", action="store_true", help="输出调试日志")
    args = ap.parse_args()

    if args.template:
        make_template()
        return 0
    if not args.file:
        ap.error("请使用 -f 指定 Excel 文件，或用 --template 生成模板")

    args.workers = max(1, min(args.workers, 50))
    args.timeout = max(3, args.timeout)
    args.retries = max(0, min(args.retries, 5))

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.FileHandler(f"save_config_{stamp}.log", encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    logging.getLogger("paramiko").setLevel(logging.CRITICAL)  # 屏蔽 paramiko 噪音
    if args.legacy_ssh:
        enable_legacy_algorithms()
        log.info("已启用旧版 SSH 算法兼容（--legacy-ssh）")

    devices, results = read_devices(args.file)
    for r in results:
        log.warning("第 %d 行被跳过: %s", r.dev.row, r.detail)
    log.info("有效设备 %d 台，参数不合法/重复 %d 台；模式: %s",
             len(devices), len(results), "仅校验" if args.check_only else "校验并保存")
    if not devices:
        results_path = write_report(results, stamp)
        log.error("没有可处理的设备，报告: %s", results_path)
        return 1

    interrupted = False
    pool = ThreadPoolExecutor(max_workers=min(args.workers, len(devices)))
    futures = {pool.submit(process_device, d, args.timeout, args.retries, args.check_only,
                           args.backup): d
               for d in devices}
    try:
        for fut in as_completed(futures):
            dev = futures[fut]
            try:
                res = fut.result()
            except Exception as e:  # noqa: BLE001  兜底：线程内未预期异常不影响其他设备
                res = Result(dev, STATUS_FAIL, "未知异常", f"{type(e).__name__}: {_short(e)}")
            results.append(res)
            lvl = logging.INFO if res.status in (STATUS_OK, STATUS_CHECK_OK) else logging.ERROR
            log.log(lvl, "[%s] %s (%s:%s) %s%s", res.status, dev.name, dev.ip, dev.port,
                    f"<{res.stage}> " if res.stage else "", res.detail)
    except KeyboardInterrupt:
        interrupted = True
        log.warning("收到中断信号，取消未开始的任务，等待运行中的任务结束…")
        pool.shutdown(wait=True, cancel_futures=True)
        for fut, dev in futures.items():
            if fut.done() and not fut.cancelled() and all(r.dev is not dev for r in results):
                try:
                    results.append(fut.result())
                except Exception:  # noqa: BLE001
                    pass
        done_devs = {id(r.dev) for r in results}
        for dev in devices:
            if id(dev) not in done_devs:
                results.append(Result(dev, STATUS_SKIP, "用户中断", "未执行"))
    finally:
        pool.shutdown(wait=True)

    report = write_report(results, stamp)
    stat = {}
    for r in results:
        stat[r.status] = stat.get(r.status, 0) + 1
    log.info("完成：%s | 报告: %s",
             "，".join(f"{k} {v}" for k, v in stat.items()), report)

    if interrupted:
        return 130
    bad = sum(v for k, v in stat.items() if k in (STATUS_FAIL, STATUS_UNSURE, STATUS_SKIP))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())