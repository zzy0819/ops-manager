"""报表服务层"""
from datetime import datetime, timedelta
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy import func
from app.extensions import db
from app.models import AttackEvent, ConfigChange, DesktopTicket, DeviceAsset


class ReportService:
    """报表业务逻辑"""

    @staticmethod
    def get_report_period(report_type='weekly', start_str='', end_str=''):
        now = datetime.now()
        if start_str and end_str:
            try:
                period_start = datetime.strptime(start_str, '%Y-%m-%d')
                period_end = datetime.strptime(end_str, '%Y-%m-%d') + timedelta(days=1)
            except ValueError:
                period_start = now - timedelta(days=7)
                period_end = now
        elif report_type == 'monthly':
            period_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            period_end = now
        else:
            period_start = now - timedelta(days=7)
            period_end = now

        period_start_display = period_start.strftime('%Y-%m-%d')
        period_end_display = (period_end - timedelta(days=1)).strftime('%Y-%m-%d')
        delta_days = (period_end - period_start).days

        return period_start, period_end, period_start_display, period_end_display, delta_days

    @staticmethod
    def get_attack_stats(period_start, period_end):
        total = AttackEvent.query.filter(
            AttackEvent.attack_time >= period_start, AttackEvent.attack_time < period_end
        ).count()

        by_type = {}
        for atype in AttackEvent.TYPE_LIST:
            by_type[atype] = AttackEvent.query.filter(
                AttackEvent.attack_time >= period_start, AttackEvent.attack_time < period_end,
                AttackEvent.attack_type == atype
            ).count()

        by_severity = {}
        for sev in AttackEvent.SEVERITY_LIST:
            by_severity[sev] = AttackEvent.query.filter(
                AttackEvent.attack_time >= period_start, AttackEvent.attack_time < period_end,
                AttackEvent.severity == sev
            ).count()

        trend_labels = []
        trend_data = []
        delta = (period_end - period_start).days
        for i in range(delta):
            day = period_start + timedelta(days=i)
            day_end = day + timedelta(days=1)
            trend_labels.append(day.strftime('%m-%d'))
            trend_data.append(AttackEvent.query.filter(
                AttackEvent.attack_time >= day, AttackEvent.attack_time < day_end
            ).count())

        return {
            'total': total,
            'by_type': by_type,
            'by_severity': by_severity,
            'trend_labels': trend_labels,
            'trend_data': trend_data
        }

    @staticmethod
    def get_config_stats(period_start, period_end):
        total = ConfigChange.query.filter(
            ConfigChange.change_time >= period_start, ConfigChange.change_time < period_end
        ).count()

        by_type = {}
        for ctype in ConfigChange.CONFIG_TYPE_LIST:
            by_type[ctype] = ConfigChange.query.filter(
                ConfigChange.change_time >= period_start, ConfigChange.change_time < period_end,
                ConfigChange.config_type == ctype
            ).count()

        by_status = {}
        for status in ['成功', '失败', '已回滚']:
            by_status[status] = ConfigChange.query.filter(
                ConfigChange.change_time >= period_start, ConfigChange.change_time < period_end,
                ConfigChange.change_status == status
            ).count()

        trend_labels = []
        trend_data = []
        delta = (period_end - period_start).days
        for i in range(delta):
            day = period_start + timedelta(days=i)
            day_end = day + timedelta(days=1)
            trend_labels.append(day.strftime('%m-%d'))
            trend_data.append(ConfigChange.query.filter(
                ConfigChange.change_time >= day, ConfigChange.change_time < day_end
            ).count())

        return {
            'total': total,
            'by_type': by_type,
            'by_status': by_status,
            'trend_labels': trend_labels,
            'trend_data': trend_data
        }

    @staticmethod
    def get_ticket_stats(period_start, period_end):
        total = DesktopTicket.query.filter(
            DesktopTicket.create_time >= period_start, DesktopTicket.create_time < period_end
        ).count()
        pending = DesktopTicket.query.filter(
            DesktopTicket.create_time >= period_start, DesktopTicket.create_time < period_end,
            DesktopTicket.status.in_(['待处理', '处理中'])
        ).count()
        resolved = DesktopTicket.query.filter(
            DesktopTicket.create_time >= period_start, DesktopTicket.create_time < period_end,
            DesktopTicket.status == '已解决'
        ).count()

        return {
            'total': total,
            'pending': pending,
            'resolved': resolved
        }

    @staticmethod
    def get_device_stats():
        total = DeviceAsset.query.count()
        by_type = db.session.query(
            DeviceAsset.device_type, func.count(DeviceAsset.id)
        ).group_by(DeviceAsset.device_type).all()

        return {
            'total': total,
            'by_type': by_type
        }

    @staticmethod
    def generate_excel_report(period_start, period_end):
        wb = openpyxl.Workbook()

        header_font = Font(name='微软雅黑', bold=True, size=12, color='FFFFFF')
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        data_font = Font(name='微软雅黑', size=10)

        def style_header(ws, row, cols):
            for col in range(1, cols + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

        def style_data_row(ws, row, cols):
            for col in range(1, cols + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')

        # 概览
        ws1 = wb.active
        ws1.title = '报告概览'
        period_label = f'{period_start.strftime("%Y-%m-%d")} ~ {(period_end - timedelta(days=1)).strftime("%Y-%m-%d")}'
        ws1.merge_cells('A1:C1')
        ws1['A1'] = f'运维管理周报/月报 - {period_label}'
        ws1['A1'].font = Font(name='微软雅黑', bold=True, size=14)

        attack_stats = ReportService.get_attack_stats(period_start, period_end)
        config_stats = ReportService.get_config_stats(period_start, period_end)
        ticket_stats = ReportService.get_ticket_stats(period_start, period_end)
        device_stats = ReportService.get_device_stats()

        overview_data = [
            ['指标', '数量', '备注'],
            ['攻击事件总数', attack_stats['total'], ''],
            ['配置变更总数', config_stats['total'], ''],
            ['工单总数', ticket_stats['total'], ''],
            ['设备资产总数', device_stats['total'], ''],
        ]
        for r, row_data in enumerate(overview_data, 2):
            for c, val in enumerate(row_data, 1):
                ws1.cell(row=r, column=c, value=val)
        style_header(ws1, 2, 3)
        for r in range(3, 3 + len(overview_data) - 1):
            style_data_row(ws1, r, 3)
        ws1.column_dimensions['A'].width = 20
        ws1.column_dimensions['B'].width = 15
        ws1.column_dimensions['C'].width = 25

        # 攻击事件分析
        ws2 = wb.create_sheet('攻击事件分析')
        ws2.merge_cells('A1:D1')
        ws2['A1'] = f'攻击事件分析 - 总计 {attack_stats["total"]} 起'
        ws2['A1'].font = Font(name='微软雅黑', bold=True, size=13)

        ws2['A3'] = '按攻击类型'
        ws2['A3'].font = Font(name='微软雅黑', bold=True, size=11)
        attack_type_headers = ['攻击类型', '数量']
        for c, h in enumerate(attack_type_headers, 1):
            ws2.cell(row=4, column=c, value=h)
        style_header(ws2, 4, 2)
        row = 5
        for atype, cnt in attack_stats['by_type'].items():
            ws2.cell(row=row, column=1, value=atype)
            ws2.cell(row=row, column=2, value=cnt)
            style_data_row(ws2, row, 2)
            row += 1

        # 配置变更分析
        ws3 = wb.create_sheet('配置变更分析')
        ws3.merge_cells('A1:D1')
        ws3['A1'] = f'配置变更分析 - 总计 {config_stats["total"]} 次'
        ws3['A1'].font = Font(name='微软雅黑', bold=True, size=13)

        ws3['A3'] = '按配置类型'
        ws3['A3'].font = Font(name='微软雅黑', bold=True, size=11)
        ctype_headers = ['配置类型', '数量']
        for c, h in enumerate(ctype_headers, 1):
            ws3.cell(row=4, column=c, value=h)
        style_header(ws3, 4, 2)
        row = 5
        for ctype, cnt in config_stats['by_type'].items():
            ws3.cell(row=row, column=1, value=ctype)
            ws3.cell(row=row, column=2, value=cnt)
            style_data_row(ws3, row, 2)
            row += 1

        # 工单与设备
        ws4 = wb.create_sheet('工单与设备')
        ws4.merge_cells('A1:C1')
        ws4['A1'] = '工单与设备资产统计'
        ws4['A1'].font = Font(name='微软雅黑', bold=True, size=13)

        ticket_data = [
            ['工单统计', '数量', ''],
            ['工单总数', ticket_stats['total'], ''],
            ['待处理/处理中', ticket_stats['pending'], ''],
            ['已解决', ticket_stats['resolved'], ''],
            ['解决率', f'{ticket_stats["resolved"] / ticket_stats["total"] * 100:.1f}%' if ticket_stats["total"] > 0 else 'N/A', ''],
        ]
        for r, row_data in enumerate(ticket_data, 2):
            for c, val in enumerate(row_data, 1):
                ws4.cell(row=r, column=c, value=val)
        style_header(ws4, 2, 3)
        for r in range(3, 3 + len(ticket_data) - 1):
            style_data_row(ws4, r, 3)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
